"""生成器使用合成生产单与合成模板；不访问客户原件或正式数据库。"""
from contextlib import contextmanager
from copy import deepcopy
import importlib.util
import json
from pathlib import Path
import sqlite3
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.pagebreak import Break

from app.production_documents import ProductionDocuments, _expanded


class SyntheticStore:
    def __init__(self, root):
        self.workspace = Path(root)
        self.path = self.workspace / "test.sqlite3"
        with self.connect() as conn:
            conn.executescript("""
                CREATE TABLE users(username TEXT,display_name TEXT,role TEXT,active INTEGER);
                INSERT INTO users VALUES('admin','测试管理员','admin',1);
                INSERT INTO users VALUES('viewer','测试查看人','viewer',1);
                CREATE TABLE audit(time TEXT,username TEXT,action TEXT,item_id TEXT,details_json TEXT);
            """)

    @contextmanager
    def connect(self):
        conn = sqlite3.connect(self.path)
        conn.row_factory = sqlite3.Row
        try:
            yield conn
            conn.commit()
        except Exception:
            conn.rollback()
            raise
        finally:
            conn.close()


def template(root, kind):
    folder = Path(root) / "private" / "production_templates"
    folder.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    sheet = book.active
    sheet.title = "合成模板"
    sheet["A1"] = "跟单 共件第件"
    sheet["A4"] = "实际检查结果留空"
    sheet["A18"] = "检查人签字："
    sheet["F2"].font = Font(name="宋体", size=11)
    sheet["F2"].alignment = Alignment(wrap_text=True)
    sheet["F2"].fill = PatternFill("solid", fgColor="FFFFCC")
    sheet.row_dimensions[6].height = 80
    sheet.column_dimensions["B"].width = 22
    sheet.merge_cells("F2:G2")
    sheet.print_area = "A1:M18" if kind == "fitter" else "A1:O13"
    sheet.print_title_rows = "1:2"
    sheet.print_title_cols = "A:A"
    sheet.row_breaks.append(Break(id=10))
    sheet.col_breaks.append(Break(id=11))
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.scale = 88
    sheet.page_margins.left = 0.3
    sheet.oddHeader.center.text = "合成页眉"
    sheet.oddFooter.right.text = "第&P页"
    sheet.print_options.horizontalCentered = True
    book.save(folder / (kind + ".xlsx"))
    book.close()


def fake_pdf(source, destination):
    from pypdf import PdfWriter
    writer = PdfWriter()
    writer.add_blank_page(width=842, height=595)
    with destination.open("wb") as handle:
        writer.write(handle)
    return 1


class ProductionDocumentTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.store = SyntheticStore(self.root)
        self.documents = ProductionDocuments(self.store)
        self.record = {
            "id": "document-1", "order_id": "ORDER-SYNTHETIC", "filename": "合成生产单.xlsx", "sha256": "synthetic-source-hash",
            "fields": {"buyer": "合成买方", "order_no": "TEST-001", "business_contact": "合成业务员", "followup": "",
                       "voltage": "220V", "plug": "", "startup": ""},
            "products": [{"sequence": "1", "name": "Synthetic Bear\n合成熊", "quantity": 2, "unit": "个",
                          "specification": "高2米1个、高3米1个", "piece_specifications": ["高2米", "高3米"],
                          "motions": "1.点头\n2.手臂摆动", "source_ref": "Sheet1!B8:H8"}], "checks": []}
        self.intake = patch.object(self.documents, "_intake", side_effect=lambda order: {"documents": [self.record], "latest": self.record})
        self.intake.start()
        self.pdf = patch("app.production_documents._convert_pdf", side_effect=fake_pdf)
        self.pdf.start()

    def tearDown(self):
        self.pdf.stop()
        self.intake.stop()
        self.temp.cleanup()

    def test_quantity_and_exact_size_expansion_with_mixed_units(self):
        self.record["products"].append({"name": "合成树", "quantity": 1, "unit": "组", "piece_specifications": ["高5米"]})
        result = self.documents.generate("ORDER-SYNTHETIC", "labels", "admin")
        self.assertEqual(result["count"], 3)
        self.assertTrue(result["pdf_ready"])
        path, name, mime = self.documents.file(result["id"], "xlsx")
        with path.open("rb") as handle:
            book = load_workbook(handle)
            self.assertIn("产品名字：合成熊\nSynthetic Bear", book.active["A1"].value)
            self.assertIn("序号：1", book.active["A1"].value)
            self.assertIn("产品尺寸：高2米", book.active["A1"].value)
            self.assertIn("产品尺寸：高3米", book.active["B1"].value)
            self.assertIn("数量：1组（按组）", book.active["C1"].value)
            self.assertEqual(book.active.max_column, 3)
            book.close()
        self.assertNotIn("files", result)

    def test_invalid_quantity_is_omitted_with_explicit_warning_and_never_defaulted_to_one(self):
        for qty in (None, 0, -2, True, 1.5, "3"):
            with self.subTest(quantity=qty):
                record = deepcopy(self.record)
                record["products"].append({"name": "异常产品", "quantity": qty, "unit": "个"})
                pieces, warnings = _expanded(record)
                self.assertEqual(len(pieces), 2)
                self.assertTrue(any("异常产品" in value and "未展开" in value for value in warnings))

    def test_conflicting_size_never_reduces_source_quantity(self):
        self.record["products"][0].update(quantity=5, specification="三种规格各1个", piece_specifications=[""] * 5)
        result = self.documents.generate("ORDER-SYNTHETIC", "labels", "admin")
        self.assertEqual(result["count"], 5)
        self.assertTrue(any("尺寸留空" in value for value in result["warnings"]))
        path, _, _ = self.documents.file(result["id"], "xlsx")
        book = load_workbook(path)
        self.assertNotIn("三种规格", book.active["A1"].value)
        self.assertIn("产品尺寸：\n", book.active["A1"].value)
        book.close()

    def test_total_expansion_limit_does_not_silently_shrink_quantities(self):
        record = deepcopy(self.record)
        record["products"] = [{"name": "大批量甲", "quantity": 6000}, {"name": "大批量乙", "quantity": 6000}]
        with self.assertRaisesRegex(ValueError, "12000.*未缩减"):
            _expanded(record)

    def test_print_setup_styles_headers_and_manual_breaks_survive_template_copy(self):
        template(self.root, "fitter")
        result = self.documents.generate("ORDER-SYNTHETIC", "fitter", "admin")
        path, _, _ = self.documents.file(result["id"], "xlsx")
        book = load_workbook(path)
        self.assertEqual(book.sheetnames, ["001", "002"])
        for sheet in book:
            self.assertEqual(str(sheet.print_area), "'" + sheet.title + "'!$A$1:$M$18")
            self.assertEqual(sheet.print_title_rows, "$1:$2")
            self.assertEqual(sheet.print_title_cols, "$A:$A")
            self.assertEqual(sheet.row_breaks.brk[0].id, 10)
            self.assertEqual(sheet.col_breaks.brk[0].id, 11)
            self.assertEqual(sheet.oddHeader.center.text, "合成页眉")
            self.assertEqual(sheet.oddFooter.right.text, "第&P页")
            self.assertEqual(sheet.page_setup.scale, 88)
            self.assertEqual(sheet.page_setup.orientation, "landscape")
            self.assertEqual(sheet.page_margins.left, 0.3)
            self.assertTrue(sheet.print_options.horizontalCentered)
            self.assertEqual(sheet.row_dimensions[6].height, 80)
            self.assertEqual(sheet.column_dimensions["B"].width, 22)
            self.assertEqual(sheet["F2"].font.name, "宋体")
            self.assertEqual(sheet["A18"].value, "检查人签字：")
            self.assertEqual(sheet["B2"].value, "合成买方")
            self.assertEqual(sheet["K2"].value, 1)
        self.assertEqual(book["001"]["I2"].value, "高2米")
        self.assertEqual(book["002"]["I2"].value, "高3米")
        book.close()

    def test_product_mapping_and_group_is_explicit(self):
        template(self.root, "product")
        source_path = self.root / "private" / "production_templates" / "product.xlsx"
        source_book = load_workbook(source_path)
        source_book.active["A1"] = "大洋艺术仿真展品工艺质量检测表 NO.共 件第 件 跟单"
        source_book.save(source_path)
        source_book.close()
        self.record["products"] = [{"name": "合成女巫", "quantity": 1, "unit": "组", "piece_specifications": ["1.5米"], "motions": "扫把动"}]
        result = self.documents.generate("ORDER-SYNTHETIC", "product", "admin")
        path, _, _ = self.documents.file(result["id"], "xlsx")
        book = load_workbook(path)
        sheet = book.active
        self.assertEqual(sheet["B2"].value, "TEST-001")
        self.assertEqual(sheet["J2"].value, "1组")
        self.assertEqual(sheet["H2"].value, "1.5米")
        self.assertEqual(sheet["B6"].value, "扫把动")
        self.assertIn("按组", sheet["A1"].value)
        self.assertTrue(sheet["A1"].value.startswith("大洋艺术仿真展品工艺质量检测表"))
        book.close()

    def test_missing_information_stays_empty_and_formula_like_source_is_text(self):
        template(self.root, "product")
        self.record["fields"].update(order_no="=HYPERLINK(\"https://example.invalid\")", business_contact="", voltage="", plug="")
        result = self.documents.generate("ORDER-SYNTHETIC", "product", "admin")
        path, _, _ = self.documents.file(result["id"], "xlsx")
        book = load_workbook(path)
        self.assertEqual(book.active["B2"].data_type, "s")
        self.assertIsNone(book.active["D2"].value)
        self.assertIsNone(book.active["L2"].value)
        book.close()

    def test_renderer_unavailable_keeps_xlsx_and_reports_actual_failure(self):
        with patch("app.production_documents._convert_pdf", side_effect=RuntimeError("未找到本机 LibreOffice")):
            result = self.documents.generate("ORDER-SYNTHETIC", "labels", "admin")
            self.assertFalse(result["pdf_ready"])
            self.assertEqual(result["pages"], 0)
            self.assertIn("LibreOffice", result["pdf_error"])
            self.assertTrue(self.documents.file(result["id"], "xlsx")[0].is_file())
            with self.assertRaisesRegex(ValueError, "LibreOffice"):
                self.documents.file(result["id"], "pdf")

    def test_download_hash_and_scope_are_checked(self):
        result = self.documents.generate("ORDER-SYNTHETIC", "labels", "admin")
        path, _, _ = self.documents.file(result["id"], "xlsx")
        original = path.read_bytes()
        path.write_bytes(b"tampered")
        with self.assertRaisesRegex(ValueError, "校验"):
            self.documents.file(result["id"], "xlsx")
        path.write_bytes(original)
        record = self.documents._get(result["id"])
        record["files"]["xlsx"]["path"] = "../../test.sqlite3"
        self.documents._save(record)
        with self.assertRaisesRegex(ValueError, "校验"):
            self.documents.file(result["id"], "xlsx")
        with self.assertRaises(KeyError):
            self.documents.file("../other", "xlsx")

    def test_history_versions_and_permissions_and_empty_rows(self):
        first = self.documents.generate("ORDER-SYNTHETIC", "labels", "admin")
        second = self.documents.generate("ORDER-SYNTHETIC", "labels", "admin")
        self.assertEqual((first["batch"], second["batch"]), (1, 2))
        self.assertNotEqual(first["filename"], second["filename"])
        status = self.documents.status("ORDER-SYNTHETIC")
        self.assertEqual(len(status["history"]), 2)
        self.assertFalse(next(item for item in status["templates"] if item["kind"] == "fitter")["ready"])
        with self.assertRaises(PermissionError):
            self.documents.generate("ORDER-SYNTHETIC", "labels", "viewer")
        self.record["products"] = [{"name": "缺数量", "quantity": None}]
        with self.assertRaisesRegex(ValueError, "未展开"):
            self.documents.generate("ORDER-SYNTHETIC", "labels", "admin")

    def test_preview_page_comes_from_pdf_and_bounds_are_enforced(self):
        result = self.documents.generate("ORDER-SYNTHETIC", "labels", "admin")
        if not (importlib.util.find_spec("pypdfium2") or importlib.util.find_spec("fitz")):
            self.skipTest("test environment has no PDF image renderer")
        preview = self.documents.preview_page(result["id"], 1)
        self.assertTrue(preview.read_bytes().startswith(b"\x89PNG"))
        self.assertEqual(self.documents.preview_page(result["id"], 1), preview)
        with self.assertRaisesRegex(ValueError, "页码"):
            self.documents.preview_page(result["id"], 0)
        with self.assertRaisesRegex(ValueError, "超出"):
            self.documents.preview_page(result["id"], 2)


if __name__ == "__main__":
    unittest.main()

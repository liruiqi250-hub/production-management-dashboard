"""生产资料：按原单数量展开、填写私有模板，并从实际 XLSX 转换打印 PDF。"""
from __future__ import annotations

from copy import copy, deepcopy
from contextlib import contextmanager
from datetime import datetime, timedelta
import hashlib
import importlib.util
import json
import math
from pathlib import Path
import re
import shutil
import subprocess
import tempfile
import threading
import uuid
import unicodedata

from app import __version__
from app.permissions import active_actor, permission_profile
from app.store import TZ, now


KINDS = {"labels": "产品标签", "fitter": "钳工质检表", "product": "产品质检表"}
MIME = {"xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "pdf": "application/pdf"}
MAX_PIECES = 10000


class ArtifactsExpired(ValueError):
    """历史仍在，但可重新生成的临时文件已按保留期清理。"""
    code = "artifacts_expired"


def _time(value):
    if isinstance(value, str):
        value = datetime.fromisoformat(value)
    if not isinstance(value, datetime):
        raise ValueError("资料使用时间无效。")
    return value.replace(tzinfo=TZ) if value.tzinfo is None else value.astimezone(TZ)


def _linked(path):
    return path.is_symlink() or bool(getattr(path, "is_junction", lambda: False)())


def _sha(path):
    digest = hashlib.sha256()
    with Path(path).open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _name_parts(value):
    """仅重排源单已有中英文；不猜译名。"""
    chinese, english = [], []
    for line in str(value or "").splitlines():
        line = line.strip()
        if line:
            (chinese if re.search(r"[\u3400-\u9fff]", line) else english).append(line)
    return " ".join(chinese), " ".join(english)


def _write(sheet, address, value):
    cell = sheet[address]
    cell.value = "" if value is None else value
    if isinstance(cell.value, str):
        # 源单文字不能变成 Excel 公式或外链。
        cell.data_type = "s"


def _copy_sheet(workbook, source):
    target = workbook.copy_worksheet(source)
    target.print_area = copy(source.print_area)
    target.print_title_rows = source.print_title_rows
    target.print_title_cols = source.print_title_cols
    target.row_breaks = deepcopy(source.row_breaks)
    target.col_breaks = deepcopy(source.col_breaks)
    target.HeaderFooter = deepcopy(source.HeaderFooter)
    target.print_options = copy(source.print_options)
    target.page_setup = copy(source.page_setup)
    target.page_margins = copy(source.page_margins)
    target.sheet_properties = deepcopy(source.sheet_properties)
    target.sheet_view.showGridLines = source.sheet_view.showGridLines
    return target


def _actions(value, slots):
    value = str(value or "").strip()
    if not value:
        return []
    parts = [part.strip() for part in re.split(r"\n\s*\n+|\n(?=\s*\d+[.、．)])", value) if part.strip()]
    if len(parts) > slots:
        parts = parts[:slots - 1] + ["\n".join(parts[slots - 1:])]
    return parts


def _column_points(sheet, column):
    width = sheet.sheet_format.defaultColWidth or 13
    for dimension in sheet.column_dimensions.values():
        if dimension.min <= column <= dimension.max and dimension.width is not None:
            width = dimension.width
            break
    return (width * 7 + 5) * 0.75


def _content_height(sheet, cell):
    last_column = cell.column
    for area in sheet.merged_cells.ranges:
        if cell.coordinate in area:
            last_column = area.max_col
            break
    width = sum(_column_points(sheet, column) for column in range(cell.column, last_column + 1)) - 6
    font_size = float(cell.font.sz or 11)
    capacity = max(1, width / (font_size * 0.53))
    lines = 0
    for line in str(cell.value or "").split("\n"):
        units = sum(2 if unicodedata.east_asian_width(char) in "WF" else 1 for char in line)
        lines += max(1, math.ceil(units / capacity))
    return math.ceil(lines * font_size * 1.22 + 5)


def _adapt_print_content(sheet, original, action_row, kind):
    """只适配生成副本中有填入内容的数据行；原私有模板保持不变。"""
    for row in (2, action_row):
        current = sheet.row_dimensions[row].height or sheet.sheet_format.defaultRowHeight or 15
        needed = max((_content_height(sheet, cell) for cell in sheet[row] if cell.value), default=current)
        sheet.row_dimensions[row].height = max(current, needed)
    # 有明确分页/打印区域的其它模板继续沿用；公司当前模板未定义打印区域。
    if original.print_area or original.row_breaks.brk or original.col_breaks.brk:
        return False
    from openpyxl.utils import get_column_letter
    sheet.print_area = f"A1:{get_column_letter(original.max_column)}{original.max_row}"
    total_height = sum(sheet.row_dimensions[row].height or sheet.sheet_format.defaultRowHeight or 15
                       for row in range(1, original.max_row + 1))
    total_width = sum(_column_points(sheet, column) for column in range(1, original.max_column + 1))
    margins = sheet.page_margins
    printable_height = 595.28 - (margins.top + margins.bottom) * 72
    printable_width = 841.89 - (margins.left + margins.right) * 72
    fitted = min(printable_height / total_height, printable_width / total_width) * 100
    # 产品模板原始正文字号10pt/缩放72%；钳工正文11pt/缩放88%。
    minimum = 70 if kind == "product" else 73
    if fitted >= minimum:
        sheet.page_setup.scale = None
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 1
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        return False
    sheet.page_setup.fitToWidth = None
    sheet.page_setup.fitToHeight = None
    sheet.sheet_properties.pageSetUpPr.fitToPage = False
    sheet.print_title_rows = "1:3"
    return True


def _expanded(record):
    pieces, warnings = [], []
    total = sum(item.get("quantity") for item in record.get("products", [])
                if type(item.get("quantity")) is int and item["quantity"] > 0)
    if total > MAX_PIECES:
        raise ValueError(f"原单合计{total}份，超过单次生成上限{MAX_PIECES}份；未缩减原数量或生成部分资料。")
    for item in record.get("products", []):
        qty = item.get("quantity")
        name = str(item.get("name") or "")
        source = str(item.get("source_ref") or "")
        if type(qty) is not int or qty < 1:
            warnings.append(f"{name or '未命名产品'}：数量无效或缺失，本行未展开（{source}）。")
            continue
        sizes = item.get("piece_specifications")
        if not isinstance(sizes, list) or len(sizes) != qty:
            sizes = [""] * qty
            warnings.append(f"{name or '未命名产品'}：逐份尺寸未确定，尺寸留空（{source}）。")
        elif item.get("specification") and not any(sizes):
            warnings.append(f"{name or '未命名产品'}：规格分配不一致或未确定，尺寸留空（{source}）。")
        for index in range(qty):
            pieces.append({**item, "piece_index": index + 1, "piece_specification": str(sizes[index] or "")})
    return pieces, warnings


def find_libreoffice():
    candidates = [shutil.which("soffice.exe"), shutil.which("soffice"),
                  r"D:\Program Files\LibreOffice\program\soffice.exe",
                  r"C:\Program Files\LibreOffice\program\soffice.exe",
                  r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"]
    try:
        import winreg
        for hive in (winreg.HKEY_CURRENT_USER, winreg.HKEY_LOCAL_MACHINE):
            for key in (r"SOFTWARE\Microsoft\Windows\CurrentVersion\App Paths\soffice.exe",
                        r"SOFTWARE\WOW6432Node\Microsoft\Windows\CurrentVersion\App Paths\soffice.exe"):
                try:
                    with winreg.OpenKey(hive, key) as handle:
                        candidates.append(winreg.QueryValueEx(handle, "")[0])
                except OSError:
                    pass
    except ImportError:
        pass
    found = next((Path(str(path).strip('"')) for path in candidates if path and Path(str(path).strip('"')).is_file()), None)
    # Windows 的 .exe 启动器可能在转换子进程结束前退出；.com 会等待并保留诊断输出。
    console = found.with_suffix(".com") if found else None
    return console if console and console.is_file() else found


def _convert_pdf(source, destination):
    renderer = find_libreoffice()
    if renderer is None:
        raise RuntimeError("未找到本机 LibreOffice，Excel 已保存；安装或配置后可重新生成实际打印预览。")
    # LibreOffice 在 Windows 对深层 profile 路径兼容性较差；使用独立短临时目录。
    with tempfile.TemporaryDirectory(prefix="production_pdf_") as profile:
        command = [str(renderer), "-env:UserInstallation=" + Path(profile).resolve().as_uri(),
                   "--headless", "--norestore", "--nodefault", "--nofirststartwizard",
                   "--convert-to", "pdf:calc_pdf_Export", "--outdir", str(destination.parent), str(source)]
        try:
            process = subprocess.run(command, stdin=subprocess.DEVNULL, capture_output=True, timeout=240,
                                     creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
        except (OSError, subprocess.TimeoutExpired) as exc:
            raise RuntimeError("本机 Excel 到 PDF 转换未完成，Excel 文件已保留。") from exc
        if process.returncode or not destination.is_file():
            detail = (process.stderr or process.stdout).decode("utf-8", "replace").strip()[:500]
            raise RuntimeError("本机 LibreOffice 未生成 PDF，Excel 文件已保留。" + (" " + detail if detail else ""))
    if destination.read_bytes()[:5] != b"%PDF-":
        raise RuntimeError("本机转换结果不是有效 PDF，Excel 文件已保留。")
    from pypdf import PdfReader
    pages = len(PdfReader(str(destination)).pages)
    if pages < 1:
        raise RuntimeError("转换后的 PDF 没有可预览页面。")
    return pages


class ProductionDocuments:
    _registry_guard = threading.Lock()
    _locks = {}
    _active = {}

    def __init__(self, store):
        self.store = store
        self.root = Path(store.workspace).resolve()
        self.base = self.root / "private" / "production_document_outputs"
        self.templates = self.root / "private" / "production_templates"
        self.protected_output_ids = lambda: set()
        self.cleanup_receipt = self.root / "private" / "production_document_cleanup" / "last.json"
        self._maintenance_guard = threading.Lock()
        self._maintenance_stop = threading.Event()
        self._maintenance_thread = None
        with store.connect() as conn:
            conn.execute("""CREATE TABLE IF NOT EXISTS documents_generated(
                id TEXT PRIMARY KEY,order_id TEXT NOT NULL,kind TEXT NOT NULL,
                batch INTEGER NOT NULL,created_at TEXT NOT NULL,data_json TEXT NOT NULL)""")

    def _lock(self, output_id):
        if not isinstance(output_id, str) or not re.fullmatch(r"[a-f0-9]{32}", output_id):
            raise KeyError("生成资料不存在。")
        key = (str(self.root).casefold(), output_id)
        with self._registry_guard:
            return key, self._locks.setdefault(key, threading.RLock())

    @contextmanager
    def hold(self, output_id):
        """可重入批次占用；HTTP调用者应覆盖验证、实际读取及响应发送。"""
        key, lock = self._lock(output_id)
        with lock:
            with self._registry_guard:
                self._active[key] = self._active.get(key, 0) + 1
            try:
                yield
            finally:
                with self._registry_guard:
                    self._active[key] -= 1
                    if not self._active[key]:
                        self._active.pop(key, None)

    @staticmethod
    def _available(record):
        if record.get("status") == "cleaned" or record.get("cleaned_at"):
            raise ArtifactsExpired("临时资料已按保留期清理，生成历史和原生产单保留，请重新生成。")

    def _touch(self, output_id):
        with self.hold(output_id):
            record = self._get(output_id)
            self._available(record)
            timestamp = now()
            try:
                if _time(record.get("last_used_at") or record["created_at"]) > _time(timestamp):
                    return
            except (ValueError, KeyError, TypeError):
                pass
            record["last_used_at"] = timestamp
            self._save(record)

    def _write_cleanup_receipt(self, receipt):
        self.cleanup_receipt.parent.mkdir(parents=True, exist_ok=True)
        temporary = self.cleanup_receipt.with_name(uuid.uuid4().hex + ".tmp")
        try:
            temporary.write_text(json.dumps(receipt, ensure_ascii=False, indent=2), encoding="utf-8")
            temporary.replace(self.cleanup_receipt)
        finally:
            temporary.unlink(missing_ok=True)

    def _cleanup_file_plan(self, record):
        """先完整核对一个批次，未知或非原样文件使整批跳过。"""
        base = self.root / "private" / "production_document_outputs"
        if self.base != base or _linked(base) or _linked(base.parent):
            raise ValueError("清理目录不是本项目固定的临时资料目录。")
        base_resolved = base.resolve()
        output_id = record["id"]
        if not re.fullmatch(r"[a-f0-9]{32}", output_id):
            raise ValueError("临时资料批次编号无效。")
        batch = base / output_id
        if _linked(batch) or not batch.is_dir() or batch.resolve().parent != base_resolved:
            raise ValueError("临时资料批次目录缺失、存在链接或范围异常。")
        infos = record.get("files")
        if not isinstance(infos, dict) or not infos:
            raise ValueError("该批次没有已登记的可再生文件。")
        previous_removed = set(record.get("cleanup_removed") or [])
        expected, plan = set(), []
        for key, info in infos.items():
            if not isinstance(key, str) or not isinstance(info, dict):
                raise ValueError("生成资料登记信息异常。")
            extension = {"xlsx": ".xlsx", "pdf": ".pdf"}.get(key)
            if re.fullmatch(r"(?:print_)?preview_[1-9]\d*", key):
                extension = ".png"
            if not extension or not isinstance(info.get("path"), str) or not re.fullmatch(r"[a-f0-9]{64}", str(info.get("sha256", ""))):
                raise ValueError("文件未列入已登记的Excel/PDF/PNG白名单。")
            relative = Path(info["path"])
            if relative.is_absolute() or ".." in relative.parts or not relative.parts or relative.parts[0] != output_id:
                raise ValueError("已登记文件路径超出该批次。")
            if len(relative.parts) != (3 if extension == ".png" else 2) or (extension == ".png" and relative.parts[1] != "previews"):
                raise ValueError("已登记文件不在允许的批次层级。")
            path = base / relative
            if path.suffix.lower() != extension or path in expected:
                raise ValueError("已登记文件类型或路径重复。")
            expected.add(path)
            if _linked(path) or _linked(path.parent) or not path.resolve().is_relative_to(batch.resolve()):
                raise ValueError("已登记文件存在链接或范围异常。")
            if not path.is_file():
                if str(relative) in previous_removed:
                    continue
                raise ValueError("已登记文件缺失，未继续清理。")
            if _sha(path) != info["sha256"]:
                raise ValueError("文件已经修改，未清理该批次。")
            plan.append((path, info["sha256"], path.stat().st_size, str(relative)))
        for entry in batch.iterdir():
            if _linked(entry):
                raise ValueError("批次中存在链接，未清理。")
            if entry.is_dir():
                if entry.name != "previews":
                    raise ValueError("批次中存在未登记目录，未清理。")
                children = list(entry.iterdir())
            else:
                children = [entry]
            for child in children:
                if _linked(child) or not child.is_file() or child not in expected:
                    raise ValueError("批次中存在未登记文件或目录，未清理。")
        return base_resolved, batch, plan

    def cleanup_expired(self, now=None, retention_days=7, dry_run=False):
        """只清理超过保留期、未占用且逐文件摘要一致的可再生输出。"""
        if type(retention_days) is not int or not 1 <= retention_days <= 3650 or type(dry_run) is not bool:
            raise ValueError("临时资料保留天数或试运行参数无效。")
        reference = _time(now if now is not None else datetime.now(TZ))
        timestamp, cutoff = reference.isoformat(timespec="seconds"), reference - timedelta(days=retention_days)
        receipt = {"status": "completed", "started_at": timestamp, "retention_days": retention_days,
                   "dry_run": dry_run, "examined": 0, "cleaned": 0, "would_clean": 0, "bytes_freed": 0,
                   "cleaned_ids": [], "skipped": [], "errors": []}
        with self.store.connect() as conn:
            candidates = [(row["id"], json.loads(row["data_json"])) for row in conn.execute("SELECT id,data_json FROM documents_generated")]
        for output_id, _ in candidates:
            receipt["examined"] += 1
            try:
                key, lock = self._lock(output_id)
            except KeyError:
                receipt["skipped"].append({"id": output_id, "reason": "批次编号无效"})
                continue
            if not lock.acquire(blocking=False):
                receipt["skipped"].append({"id": output_id, "reason": "正在生成、读取或打印"})
                continue
            try:
                with self._registry_guard:
                    active = self._active.get(key, 0)
                if active or output_id in set(self.protected_output_ids()):
                    receipt["skipped"].append({"id": output_id, "reason": "正在使用或处于打印队列"})
                    continue
                record = self._get(output_id)
                if record.get("id") != output_id:
                    raise ValueError("批次登记编号不一致。")
                if record.get("status") == "cleaned":
                    continue
                if record.get("status") not in {"ready", "failed", "cleanup_failed"}:
                    receipt["skipped"].append({"id": output_id, "reason": "仍在生成或状态未明确"})
                    continue
                try:
                    last_used = _time(record.get("last_used_at") or record["created_at"])
                except (TypeError, ValueError, KeyError):
                    receipt["skipped"].append({"id": output_id, "reason": "最后使用时间无效"})
                    continue
                if last_used > cutoff:
                    continue
                try:
                    base, batch, plan = self._cleanup_file_plan(record)
                except (ValueError, OSError, KeyError, TypeError) as exc:
                    receipt["skipped"].append({"id": output_id, "reason": str(exc)})
                    continue
                if dry_run:
                    receipt["would_clean"] += 1
                    continue
                removed = list(record.get("cleanup_removed") or [])
                try:
                    for path, expected_sha, size, relative in plan:
                        # 实际unlink前再次核对绝对目录、链接和内容，绝不递归删除。
                        if self.base.resolve() != base or _linked(self.base) or _linked(batch) or _linked(path.parent) or _linked(path) or path.resolve().parent not in {batch.resolve(), (batch / "previews").resolve()} or not path.resolve().is_relative_to(base):
                            raise ValueError("删除前绝对路径范围核验失败。")
                        if _sha(path) != expected_sha:
                            raise ValueError("删除前发现文件改变，停止该批次清理。")
                        path.unlink()
                        removed.append(relative)
                        receipt["bytes_freed"] += size
                    for directory in (batch / "previews", batch):
                        if directory.is_dir() and not _linked(directory) and not _linked(batch) and self.base.resolve() == base and directory.resolve().is_relative_to(base):
                            try:
                                directory.rmdir()  # 非空或Windows占用时只保留空目录，不强制删除。
                            except OSError:
                                pass
                    record.update(status="cleaned", cleaned_at=timestamp, last_used_at=last_used.isoformat(timespec="seconds"),
                                  pdf_ready=False, preview_ready=False, cleanup_removed=removed)
                    record.pop("cleanup_error", None)
                    self._save(record)
                    receipt["cleaned"] += 1
                    receipt["cleaned_ids"].append(output_id)
                except (OSError, ValueError) as exc:
                    record["cleanup_error"] = str(exc)
                    record["cleanup_removed"] = removed
                    if removed:
                        record["status"] = "cleanup_failed"
                    self._save(record)
                    receipt["errors"].append({"id": output_id, "message": "文件占用、权限或校验异常，未完成该批次清理：" + str(exc)})
            except Exception as exc:
                receipt["errors"].append({"id": output_id, "message": "未清理该批次：" + str(exc)})
            finally:
                lock.release()
        if receipt["errors"]:
            receipt["status"] = "partial"
        receipt["finished_at"] = datetime.now(TZ).isoformat(timespec="seconds")
        self._write_cleanup_receipt(receipt)
        return receipt

    def start_maintenance(self, *, interval_seconds=6 * 60 * 60):
        """服务生命周期调用：后台立即清理一次，随后每六小时检查。"""
        if isinstance(interval_seconds, bool) or not isinstance(interval_seconds, (int, float)) or interval_seconds <= 0:
            raise ValueError("资料维护间隔无效。")
        with self._maintenance_guard:
            if self._maintenance_thread and self._maintenance_thread.is_alive():
                return False
            self._maintenance_stop.clear()
            def maintain():
                while not self._maintenance_stop.is_set():
                    try:
                        self.cleanup_expired()
                    except Exception as exc:
                        try:
                            self._write_cleanup_receipt({"status": "failed", "finished_at": now(), "errors": [{"message": str(exc)}]})
                        except OSError:
                            pass
                    if self._maintenance_stop.wait(interval_seconds):
                        break
            self._maintenance_thread = threading.Thread(target=maintain, name="production-document-maintenance", daemon=True)
            self._maintenance_thread.start()
            return True

    def stop_maintenance(self, *, timeout=5):
        self._maintenance_stop.set()
        thread = self._maintenance_thread
        if thread and thread is not threading.current_thread():
            thread.join(timeout=timeout)
        return not thread or not thread.is_alive()

    def _intake(self, order_id):
        from app.production_order_intake import ProductionOrderIntake
        return ProductionOrderIntake(self.store).details(order_id)

    def _template(self, kind):
        path = self.templates / (kind + ".xlsx")
        if not path.is_file():
            raise ValueError(f"本机缺少{KINDS[kind]}模板，请提供公司正式模板。")
        if path.resolve().parent != self.templates.resolve():
            raise ValueError("模板路径校验失败。")
        manifest = self.templates / "manifest.json"
        if manifest.is_file():
            data = json.loads(manifest.read_text("utf-8-sig"))
            entry = next((entry for entry in data.get("templates", []) if entry.get("kind") == kind), None)
            if entry and entry.get("template_sha256") and _sha(path) != entry["template_sha256"]:
                raise ValueError(f"{KINDS[kind]}模板校验失败，请重新核对模板版本。")
        return path

    @staticmethod
    def _public(record):
        result = {key: value for key, value in record.items() if key != "files"}
        result.setdefault("last_used_at", record.get("created_at", ""))
        result["artifacts_available"] = record.get("status") != "cleaned" and not record.get("cleaned_at")
        return result

    def status(self, order_id):
        intake = self._intake(order_id)
        templates = [{"kind": "labels", "name": KINDS["labels"], "ready": True}]
        for kind in ("fitter", "product"):
            item = {"kind": kind, "name": KINDS[kind], "ready": False}
            try:
                self._template(kind)
                item["ready"] = True
            except (ValueError, OSError, json.JSONDecodeError) as exc:
                item["reason"] = str(exc)
            templates.append(item)
        with self.store.connect() as conn:
            history = [self._public(json.loads(row["data_json"])) for row in conn.execute(
                "SELECT data_json FROM documents_generated WHERE order_id=? ORDER BY batch DESC", (order_id,))]
        return {"templates": templates, "latest": intake.get("latest"), "history": history}

    def _get(self, output_id):
        if not isinstance(output_id, str) or not re.fullmatch(r"[a-f0-9]{32}", output_id):
            raise KeyError("生成资料不存在。")
        with self.store.connect() as conn:
            row = conn.execute("SELECT data_json FROM documents_generated WHERE id=?", (output_id,)).fetchone()
        if not row:
            raise KeyError("生成资料不存在。")
        return json.loads(row["data_json"])

    def _save(self, record):
        with self.hold(record["id"]), self.store.connect() as conn:
            previous = conn.execute("SELECT data_json FROM documents_generated WHERE id=?", (record["id"],)).fetchone()
            if previous:
                current = json.loads(previous["data_json"])
                if current.get("status") == "cleaned" and record.get("status") != "cleaned":
                    raise ArtifactsExpired("临时资料已清理，请重新生成。")
                try:
                    if _time(current.get("last_used_at") or current["created_at"]) > _time(record.get("last_used_at") or record["created_at"]):
                        record["last_used_at"] = current.get("last_used_at") or current["created_at"]
                except (ValueError, KeyError, TypeError):
                    pass
            conn.execute("UPDATE documents_generated SET data_json=? WHERE id=?",
                         (json.dumps(record, ensure_ascii=False), record["id"]))

    def _record_file(self, path):
        relative = path.resolve().relative_to(self.base.resolve())
        return {"path": str(relative), "sha256": _sha(path)}

    def _verified(self, info):
        path = (self.base / info["path"]).resolve()
        if not path.is_relative_to(self.base.resolve()) or not path.is_file() or _sha(path) != info["sha256"]:
            raise ValueError("生成资料文件缺失或校验失败，请重新生成。")
        return path

    def generate(self, order_id, kind, actor, document_id=None):
        if not isinstance(kind, str) or kind not in KINDS:
            raise ValueError("资料类型无效。")
        username = actor.get("username") if isinstance(actor, dict) else actor
        with self.store.connect() as conn:
            user = active_actor(conn, username)
            if not permission_profile(user, self.root)["can_write"]:
                raise PermissionError("该账号仅有查看权限，不能生成资料。")
        intake = self._intake(order_id)
        source = (next((entry for entry in intake.get("documents", []) if entry["id"] == document_id), None)
                  if document_id else intake.get("latest"))
        if not source:
            raise ValueError("这张生产单尚未关联可读取的原件，请先导入生产单。")
        pieces, warnings = _expanded(source)
        warnings = [str(check.get("message") or "") for check in source.get("checks", []) if check.get("message")] + warnings
        if not pieces:
            raise ValueError("没有可展开的合法产品数量；" + "；".join(warnings))
        output_id, timestamp = uuid.uuid4().hex, now()
        with self.store.connect() as conn:
            conn.execute("BEGIN IMMEDIATE")
            batch = conn.execute("SELECT coalesce(max(batch),0)+1 FROM documents_generated WHERE order_id=?", (order_id,)).fetchone()[0]
            stem = re.sub(r'[\\/:*?"<>|\x00-\x1f]', "_", Path(str(source.get("filename") or "生产单")).stem).strip(" .")[:110]
            filename = (stem or "生产单") + "_" + KINDS[kind] + f"_V{batch:03d}.xlsx"
            result = {"id": output_id, "order_id": order_id, "document_id": source["id"], "kind": kind,
                      "filename": filename, "batch": batch, "version": __version__, "created_at": timestamp,
                      "count": len(pieces), "warnings": list(dict.fromkeys(warnings)), "pdf_ready": False,
                      "preview_ready": False, "pages": 0, "status": "generating", "files": {},
                      "last_used_at": timestamp,
                      "source_sha256": source.get("sha256", ""), "created_by": username,
                      "totals_by_unit": source.get("totals_by_unit", {})}
            conn.execute("INSERT INTO documents_generated VALUES(?,?,?,?,?,?)",
                         (output_id, order_id, kind, batch, timestamp, json.dumps(result, ensure_ascii=False)))
        with self.hold(output_id):
            return self._generate_artifacts(result, source, pieces, kind, username)

    def _generate_artifacts(self, result, source, pieces, kind, username):
        output_id, timestamp, order_id = result["id"], result["created_at"], result["order_id"]
        folder = self.base / output_id
        target = folder / result["filename"]
        try:
            folder.mkdir(parents=True, exist_ok=False)
            if kind == "labels":
                self._labels(source, pieces, target)
            else:
                extra = self._inspection(kind, source, pieces, target)
                result["warnings"] = list(dict.fromkeys(result["warnings"] + extra))
            result["files"]["xlsx"] = self._record_file(target)
            result["status"] = "ready"
            self._save(result)
            self.render_pdf(output_id)
            self._touch(output_id)
            result = self._get(output_id)
            with self.store.connect() as conn:
                conn.execute("INSERT INTO audit(time,username,action,item_id,details_json) VALUES(?,?,?,?,?)",
                             (timestamp, username, "generate_production_documents", order_id,
                              json.dumps(self._public(result), ensure_ascii=False)))
            return self._public(result)
        except Exception as exc:
            result["status"] = "failed"
            result["error"] = str(exc)
            self._save(result)
            raise

    def _labels(self, record, pieces, target):
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, Side
        from openpyxl.worksheet.page import PageMargins
        from openpyxl.worksheet.pagebreak import Break
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "标签"
        sheet.sheet_view.showGridLines = False
        for column in "ABC":
            sheet.column_dimensions[column].width = 34
        border = Border(*[Side(style="thin", color="808080") for _ in range(4)])
        buyer = str(record.get("fields", {}).get("buyer") or "")
        for index, piece in enumerate(pieces):
            row, column = index // 3 + 1, index % 3 + 1
            chinese, english = _name_parts(piece.get("name"))
            sequence = str(piece.get("sequence") or "")
            content = f"买方信息：{buyer}\n序号：{sequence}\n产品名字：{chinese}\n{english}\n产品尺寸：{piece['piece_specification']}"
            unit = str(piece.get("unit") or "")
            if unit == "组":
                content += "\n数量：1组（按组）"
            elif unit:
                content += "\n数量：1" + unit
            cell = sheet.cell(row, column)
            _write(sheet, cell.coordinate, content)
            cell.font = Font(name="宋体", size=11)
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            cell.border = copy(border)
            sheet.row_dimensions[row].height = 185
        sheet.page_setup.orientation = "portrait"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_margins = PageMargins(left=0.25, right=0.25, top=0.25, bottom=0.25, header=0, footer=0)
        sheet.print_area = f"A1:C{(len(pieces) + 2) // 3}"
        for row in range(5, (len(pieces) + 2) // 3 + 1, 4):
            sheet.row_breaks.append(Break(id=row - 1))
        workbook.save(target)
        workbook.close()

    def _inspection(self, kind, record, pieces, target):
        from openpyxl import load_workbook
        workbook = load_workbook(self._template(kind), keep_links=False)
        original = workbook.active
        fields = record.get("fields", {})
        multipage = []
        for index, piece in enumerate(pieces, 1):
            sheet = _copy_sheet(workbook, original)
            sheet.title = f"{index:03d}"
            unit = str(piece.get("unit") or "")
            suffix = "（按组）" if unit == "组" else ""
            title = re.split(r"NO[.．]", str(original["A1"].value or ""), maxsplit=1, flags=re.I)[0]
            _write(sheet, "A1", title + f"NO.共{len(pieces)}份 第{index}份  跟单 {fields.get('followup') or ''}{suffix}")
            chinese, english = _name_parts(piece.get("name"))
            name = "\n".join(value for value in (english, chinese) if value)
            voltage = "\n".join(line for key in ("voltage", "plug") for line in str(fields.get(key) or "").splitlines() if line.strip())
            specification = re.sub(r"\n\s*\n+", "\n", piece["piece_specification"])
            mapping = ({"B2": fields.get("buyer"), "D2": fields.get("business_contact"), "F2": name,
                        "I2": specification, "K2": "1组" if unit == "组" else 1, "M2": voltage}
                       if kind == "fitter" else
                       {"B2": fields.get("order_no"), "D2": fields.get("business_contact"), "F2": name,
                        "H2": specification, "J2": "1组" if unit == "组" else 1, "L2": voltage})
            for address, value in mapping.items():
                _write(sheet, address, value)
            row, slots = (15, 12) if kind == "fitter" else (6, 7)
            for column in range(2, slots + 2):
                _write(sheet, f"{chr(64 + column)}{row}", "")
            for column, action in enumerate(_actions(piece.get("motions"), slots), 2):
                _write(sheet, f"{chr(64 + column)}{row}", action)
            if _adapt_print_content(sheet, original, row, kind):
                multipage.append(str(index))
        workbook.remove(original)
        workbook.active = 0
        workbook.save(target)
        workbook.close()
        return (["为保留可读字号，第" + "、".join(multipage) + "份采用连续跨页打印并重复表头；实际页数以PDF预览为准。"] if multipage else [])

    def render_pdf(self, output_id):
        with self.hold(output_id):
            return self._render_pdf(output_id)

    def _render_pdf(self, output_id):
        result = self._get(output_id)
        self._available(result)
        if result.get("pdf_ready"):
            self._verified(result["files"]["pdf"])
            return self._public(result)
        xlsx = self._verified(result["files"]["xlsx"])
        pdf = xlsx.with_suffix(".pdf")
        try:
            pages = _convert_pdf(xlsx, pdf)
            result["files"]["pdf"] = self._record_file(pdf)
            result.update(pdf_ready=True, pages=pages, preview_ready=bool(importlib.util.find_spec("pypdfium2") or importlib.util.find_spec("fitz")))
            result.pop("pdf_error", None)
            if not result["preview_ready"]:
                result["preview_error"] = "PDF 已生成，缺少本机 PDF 图片渲染组件，暂不能显示逐页图片。"
        except Exception as exc:
            result["pdf_error"] = str(exc)
        self._save(result)
        return self._public(result)

    def file(self, output_id, format):
        with self.hold(output_id):
            result = self._file(output_id, format)
            self._touch(output_id)
            return result

    def _file(self, output_id, format):
        if not isinstance(format, str) or format not in MIME:
            raise ValueError("文件格式无效。")
        result = self._get(output_id)
        self._available(result)
        if format == "pdf" and not result.get("pdf_ready"):
            self.render_pdf(output_id)
            result = self._get(output_id)
        if format not in result.get("files", {}):
            raise ValueError(result.get("pdf_error") or "文件尚未生成。")
        path = self._verified(result["files"][format])
        return path, Path(result["filename"]).with_suffix("." + format).name, MIME[format]

    def preview_page(self, output_id, page, *, print_quality=False):
        with self.hold(output_id):
            path = self._preview_page(output_id, page, print_quality=print_quality)
            self._touch(output_id)
            return path

    def _preview_page(self, output_id, page, *, print_quality=False):
        if type(page) is not int or page < 1:
            raise ValueError("预览页码无效。")
        if type(print_quality) is not bool:
            raise ValueError("预览质量参数无效。")
        self.render_pdf(output_id)
        result = self._get(output_id)
        if not result.get("pdf_ready"):
            raise ValueError(result.get("pdf_error") or "PDF 预览未生成。")
        if page > result["pages"]:
            raise ValueError("预览页码超出实际 PDF 页数。")
        key = ("print_preview_" if print_quality else "preview_") + str(page)
        if key in result["files"]:
            return self._verified(result["files"][key])
        pdf = self._verified(result["files"]["pdf"])
        folder = self.base / output_id / "previews"
        folder.mkdir(exist_ok=True)
        path = folder / (f"print-page-{page}.png" if print_quality else f"page-{page}.png")
        scale = 2.5 if print_quality else 1.5
        try:
            import pypdfium2 as pdfium
        except ImportError:
            pdfium = None
        if pdfium is not None:
            document = pdfium.PdfDocument(str(pdf))
            try:
                pdf_page = document[page - 1]
                try:
                    bitmap = pdf_page.render(scale=scale)
                    try:
                        image = bitmap.to_pil()
                        try:
                            image.save(path, "PNG")
                        finally:
                            image.close()
                    finally:
                        bitmap.close()
                finally:
                    pdf_page.close()
            finally:
                document.close()
        else:
            try:
                import fitz
            except ImportError as exc:
                raise ValueError("PDF 已生成，缺少本机逐页图片预览组件。") from exc
            with fitz.open(pdf) as document:
                document[page - 1].get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False).save(path)
        result["files"][key] = self._record_file(path)
        self._save(result)
        return path
